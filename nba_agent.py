import sys
sys.stdout.reconfigure(encoding="utf-8")

import anthropic
import json
import requests
from datetime import datetime
from docx import Document

client = anthropic.Anthropic()  # reads ANTHROPIC_API_KEY from environment


def get_nba_scores():
    """Fetch today's NBA scores from ESPN's free public API (no key required)."""
    url = "https://site.api.espn.com/apis/site/v2/sports/basketball/nba/scoreboard"
    response = requests.get(url, timeout=10)
    response.raise_for_status()
    data = response.json()

    games = []
    for event in data.get("events", []):
        competition = event["competitions"][0]
        competitors = competition["competitors"]

        # ESPN orders home/away by homeAway field
        home = next((c for c in competitors if c.get("homeAway") == "home"), competitors[0])
        away = next((c for c in competitors if c.get("homeAway") == "away"), competitors[1])

        games.append({
            "home_team": home["team"]["displayName"],
            "away_team": away["team"]["displayName"],
            "home_score": home.get("score", "—"),
            "away_score": away.get("score", "—"),
            "status": competition["status"]["type"]["description"],
        })

    return games


def save_to_word_doc(scores_data: list, date_str: str) -> str:
    """Write scores to a .docx file and return the saved filename."""
    doc = Document()
    doc.add_heading(f"NBA Scores — {date_str}", level=0)

    if not scores_data:
        doc.add_paragraph("No games scheduled today.")
    else:
        for game in scores_data:
            doc.add_heading(
                f"{game['away_team']}  vs  {game['home_team']}", level=2
            )
            doc.add_paragraph(
                f"Score:  {game['away_team']} {game['away_score']}  –  "
                f"{game['home_team']} {game['home_score']}"
            )
            doc.add_paragraph(f"Status: {game['status']}")

    safe_date = date_str.replace("/", "-").replace(" ", "_").replace(",", "")
    filename = f"NBA_Scores_{safe_date}.docx"
    doc.save(filename)
    return filename


# --- Tool definitions sent to Claude ---
TOOLS = [
    {
        "name": "get_nba_scores",
        "description": "Fetches today's NBA game scores and statuses from ESPN.",
        "input_schema": {
            "type": "object",
            "properties": {},
            "required": [],
        },
    },
    {
        "name": "save_to_word_doc",
        "description": "Saves a list of NBA game scores to a Word (.docx) document.",
        "input_schema": {
            "type": "object",
            "properties": {
                "scores_data": {
                    "type": "array",
                    "description": "List of game objects returned by get_nba_scores.",
                    "items": {"type": "object"},
                },
                "date_str": {
                    "type": "string",
                    "description": "Human-readable date string, e.g. 'May 18, 2026'.",
                },
            },
            "required": ["scores_data", "date_str"],
        },
    },
]


def handle_tool(name: str, inputs: dict):
    if name == "get_nba_scores":
        return get_nba_scores()
    if name == "save_to_word_doc":
        return save_to_word_doc(inputs["scores_data"], inputs["date_str"])
    raise ValueError(f"Unknown tool: {name}")


def run_agent():
    today = datetime.now().strftime("%B %d, %Y")
    messages = [
        {
            "role": "user",
            "content": (
                f"Today is {today}. Please fetch today's NBA scores and save "
                "them to a Word document. Report what you find."
            ),
        }
    ]

    print("NBA Scores Agent starting...\n")

    while True:
        response = client.messages.create(
            model="claude-opus-4-5",
            max_tokens=4096,
            tools=TOOLS,
            messages=messages,
        )

        # Agent finished — print final message
        if response.stop_reason == "end_turn":
            for block in response.content:
                if hasattr(block, "text"):
                    print(f"Agent: {block.text}")
            break

        # Agent wants to call a tool
        if response.stop_reason == "tool_use":
            messages.append({"role": "assistant", "content": response.content})

            tool_results = []
            for block in response.content:
                if block.type == "tool_use":
                    print(f"  >> Calling tool: {block.name}")
                    result = handle_tool(block.name, block.input)
                    display = json.dumps(result, indent=2) if isinstance(result, (list, dict)) else result
                    print(f"  << Result: {display}\n")

                    tool_results.append({
                        "type": "tool_result",
                        "tool_use_id": block.id,
                        "content": json.dumps(result) if isinstance(result, (list, dict)) else str(result),
                    })

            messages.append({"role": "user", "content": tool_results})


if __name__ == "__main__":
    run_agent()
